"""Build slides/paper15_team12_LouShu2017.pptx: the same deck as deck.tex, editable in PowerPoint and Google Slides.

Every number comes from slides/gen/values.tex (written by src/build_docs.py). Speaker notes come from an optional
speech file in Markdown (one "**Slide N — ...**" block per slide). Usage:
    python -I slides/build_pptx.py [--speech path/to/speech.md]
"""
import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent

HSE = RGBColor(0x0E, 0x2D, 0x69)
SOFT = RGBColor(0xA1, 0xB8, 0xE1)
GRAY = RGBColor(0xD1, 0xD1, 0xD1)
ROSE = RGBColor(0xBF, 0x83, 0x78)
GREEN = RGBColor(0x71, 0xA3, 0x82)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x50, 0x50, 0x50)
FONT = "Arial"


def load_values(path: Path) -> dict:
    vals = {}
    for m in re.finditer(r"\\csname v@(.+?)\\endcsname\{(.*)\}\s*$", path.read_text(encoding="utf-8"), re.M):
        v = m.group(2)
        for a, b in [("\\%", "%"), ("\\_", "_"), ("\\&", "&"), ("$", ""), ("\\,", " "), ("{", ""), ("}", "")]:
            v = v.replace(a, b)
        vals[m.group(1)] = v
    text = path.read_text(encoding="utf-8")
    for name in ["basedir", "teamnames", "repourl"]:
        m = re.search(r"\\def\\" + name + r"\{(.*)\}", text)
        vals[name] = m.group(1) if m else ""
    return vals


def runs(paragraph, text, size, color=HSE, bold=False, italic=False):
    """Add text to a paragraph; _{x} becomes a subscript run."""
    for i, part in enumerate(re.split(r"_\{(.*?)\}", text)):
        if not part:
            continue
        r = paragraph.add_run()
        r.text = part
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
        f.color.rgb = color
        if i % 2:
            r._r.get_or_add_rPr().set("baseline", "-25000")


def text(slide, x, y, w, h, lines, size, color=HSE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.0, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if isinstance(line, tuple):
            runs(p, line[0], line[1], line[2] if len(line) > 2 else color, line[3] if len(line) > 3 else bold)
        else:
            runs(p, line, size, color, bold)
    return tb


def box(slide, x, y, w, h, label, fill, size=12, color=HSE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.12
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        runs(p, line, size, color)
    return s


def tri_down(slide, cx, y, w=0.26, h=0.19):
    s = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(cx - w / 2), Inches(y), Inches(w), Inches(h))
    s.rotation = 180
    s.fill.solid()
    s.fill.fore_color.rgb = HSE
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def bar(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def picture(slide, path, x=None, y=None, w=None, h=None, center_x=None):
    pic = slide.shapes.add_picture(str(path), Inches(x or 0), Inches(y or 0),
                                   Inches(w) if w else None, Inches(h) if h else None)
    if center_x is not None:
        pic.left = int(Inches(center_x) - pic.width / 2)
    return pic


class Deck:
    def __init__(self, v):
        self.v = v
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(10), Inches(5.625)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0
        self.notes = {}

    def slide(self, title=None, header=True):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        if header:
            picture(s, HERE / "hse/logo.png", 0.28, 0.22, h=0.39)
            text(s, 0.75, 0.28, 3, 0.3, ["Price Impact or Trading Volume", "(Lou, Shu, 2017)"], 7)
            text(s, 9.0, 0.26, 0.69, 0.3, str(self.n), 14, align=PP_ALIGN.RIGHT)
        if title:
            text(s, 0.47, 0.82, 9.0, 0.5, title, 24, name="Title")
        if self.n in self.notes:
            s.notes_slide.notes_text_frame.text = self.notes[self.n]
        return s

    def sentence(self, s, line, y, x=0.47, w=9.06, align=PP_ALIGN.CENTER):
        text(s, x, y, w, 0.6, line, 18, align=align)

    def stat(self, s, x, y, w, value, label, size=44, align=PP_ALIGN.CENTER):
        top = y + 0.75 - size * 1.2 / 72
        text(s, x, top, w, size * 1.2 / 72, value, size, align=align)
        text(s, x, y + 0.8, w, 0.6, label.split("\n"), 12, align=align)


def load_notes(path):
    notes = {}
    if not path or not Path(path).exists():
        return notes
    md = Path(path).read_text(encoding="utf-8")
    for m in re.finditer(r"^\*\*Slide (\d+)[^\n]*\n(.*?)(?=^\*\*Slide |^## |^---|\Z)", md, re.S | re.M):
        notes[int(m.group(1))] = m.group(2).strip()
    return notes


def build(v, notes, out):
    d = Deck(v)
    d.notes = notes
    base = ROOT / v["basedir"].replace("../", "")

    s = d.slide(header=False)
    picture(s, HERE / "hse/ornament.png", h=5.625).left = int(Inches(10) - Inches(5.625) * 1920 / 1080)
    picture(s, HERE / "hse/logo_full.png", 0.47, 0.41, h=0.72)
    text(s, 1.31, 0.47, 4.5, 0.6, ["Faculty of Computer Science", "Data Science and Business Analytics"], 12)
    text(s, 1.31, 0.95, 4.5, 0.3, "Research seminar «Data analysis in finance»", 9)
    text(s, 0.47, 1.44, 8.4, 1.0, ["Price Impact or Trading Volume:", "Why Is the Amihud (2002) Measure Priced?"], 26,
         spacing=0.95, name="Title")
    text(s, 0.47, 2.6, 8.0, 0.3, "X. Lou, T. Shu, Review of Financial Studies, 2017 · replication on Binance", 13)
    text(s, 0.47, 3.4, 8.0, 0.5, "Group 12: " + v["teamnames"], 12)
    text(s, 0.47, 4.62, 3, 0.3, "Moscow 2026", 9)

    s = d.slide("Problem")
    text(s, 0.47, 1.75, 4.2, 1.0, ["Illiquid stocks earn a premium.", "Is it for price impact, or just for low volume?"],
         18, spacing=1.05)
    picture(s, HERE / "hse/eq/amihud.png", 0.47, 2.95, w=684 / 600 * 1.9)
    for cx, head, mid, midfill, low in [(6.5, "Standard reading", "price impact", ROSE, "illiquidity\npremium"),
                                        (8.3, "Lou & Shu", "trading volume", GREEN, "low-volume\npremium")]:
        text(s, cx - 0.85, 1.35, 1.7, 0.25, head, 10, align=PP_ALIGN.CENTER)
        box(s, cx - 0.8, 1.7, 1.6, 0.52, "Amihud A", GRAY)
        tri_down(s, cx, 2.3)
        box(s, cx - 0.8, 2.57, 1.6, 0.52, mid, midfill)
        tri_down(s, cx, 3.17)
        box(s, cx - 0.8, 3.44, 1.6, 0.52, low, GRAY)

    s = d.slide("Amihud is almost only volume")
    d.stat(s, 0.8, 2.05, 3.8, v["pvu.corrAAC.paper"], "corr(A, A_{C}) · US stocks, paper", 54)
    d.stat(s, 5.4, 2.05, 3.8, v["pvu.corrAAC"], "corr(A, A_{C}) · Binance, us", 54)
    d.sentence(s, "A_{C} keeps only 1/DVol, and it explains almost all of A.", 4.15)

    s = d.slide("Method: split Amihud in two")
    picture(s, HERE / "hse/eq/split.png", 0.47, 2.02, w=1340 / 600 * 1.8)
    text(s, 0.47, 2.75, 4.0, 1.0, "If only the volume part is priced, the premium is about volume.", 18, spacing=1.05)
    box(s, 5.6, 1.55, 3.4, 0.5, "ln A", HSE, color=WHITE)
    tri_down(s, 7.3, 2.12)
    for cx, lab, fill, q in [(6.4, "ln A_{C}\nvolume", GREEN, "priced?\n“volume”"),
                             (8.2, "Res.ln A\nprice impact", ROSE, "priced?\n“price impact”")]:
        box(s, cx - 0.8, 2.4, 1.6, 0.6, lab, fill)
        tri_down(s, cx, 3.08)
        box(s, cx - 0.8, 3.35, 1.6, 0.6, q, GRAY)

    s = d.slide("Test design")
    for x, fill, top, lab in [(0.47, ROSE, "month t − 2", "MEASURE"), (3.55, GRAY, "month t − 1", "SKIP"),
                              (6.63, HSE, "month t", "RETURN")]:
        bar(s, x, 1.75, 2.9, 0.1, fill)
        text(s, x, 1.95, 2.9, 0.6, [(top, 10), (lab, 16)], 10, align=PP_ALIGN.CENTER)
    picture(s, HERE / "hse/eq/fm.png", y=2.95, w=1158 / 600 * 2.2, center_x=5.0)
    d.sentence(s, "Quintile sorts and Fama–MacBeth, as in the paper.", 4.20)

    s = d.slide("The paper's result")
    for x, lab, fill, key in [(0.6, "H-L on A", GRAY, "pvu.HLArawCoef.paper"),
                              (3.6, "H-L on A_{C}", GREEN, "pvu.HLACrawCoef.paper"),
                              (6.6, "H-L on Res.A", ROSE, "pvu.HLResArawCoef.paper")]:
        box(s, x + 0.65, 1.90, 1.6, 0.42, lab, fill)
        d.stat(s, x, 2.4, 2.9, v[key], "% per month", 26)
    d.sentence(s, "Volume is priced, the rest is not: the premium is not for price impact.", 4.25)

    s = d.slide("First, simulations")
    for x, lab, fill, key, cap in [(0.6, "World V\nvolume is priced", GREEN, "sim.V.fm2", "runs find A_{C} priced"),
                                   (3.6, "World P\nprice impact is priced", ROSE, "sim.P.fm3", "runs find Res.ln A priced"),
                                   (6.6, "World 0\nnothing is priced", GRAY, "sim.0.fm3", "false positives (target 5%)")]:
        box(s, x + 0.35, 1.85, 2.2, 0.55, lab, fill)
        d.stat(s, x, 2.75, 2.9, v[key], cap, 44)
    d.sentence(s, "The pipeline tells the hypotheses apart before we touch real data.", 4.30)

    s = d.slide("Dataset overview")
    x0, w = 0.47, 9.0
    pm = w / 93
    bar(s, x0 + 57 * pm, 1.67, 36 * pm, 0.07, ROSE)
    text(s, x0 + 57 * pm, 1.47, 3.5, 0.2, "5-min λ window, %s months" % v["hf.months"], 9, ROSE)
    bar(s, x0, 1.87, 46 * pm, 0.1, SOFT)
    bar(s, x0 + 46 * pm, 1.87, 47 * pm, 0.1, HSE)
    bar(s, x0 + 26 * pm - 0.012, 1.67, 0.025, 0.2, ROSE)
    text(s, x0 + 26 * pm - 0.6, 1.47, 1.2, 0.2, "March 2021", 9, ROSE, align=PP_ALIGN.CENTER)
    text(s, x0, 2.05, 46 * pm, 0.5, [("2019-01 – 2022-10", 9), ("FIRST HALF", 14)], 9, align=PP_ALIGN.CENTER)
    text(s, x0 + 46 * pm, 2.05, 47 * pm, 0.5, [("2022-11 – 2026-09", 9), ("SECOND HALF", 14)], 9, align=PP_ALIGN.CENTER)
    text(s, 0.47, 2.85, 9.06, 0.45, "Binance spot · USDT pairs", 26, align=PP_ALIGN.CENTER)
    for x, key, lab in [(0.47, "n.coins", "coins"), (2.67, "n.months", "months"), (4.87, "n.coinmonths", "coin-months")]:
        d.stat(s, x, 3.45, 2.1, v[key], lab, 40, align=PP_ALIGN.LEFT)
    d.stat(s, 7.3, 3.45, 2.3, "3", "signals: A, A_{C}, Res.A", 40, align=PP_ALIGN.LEFT)

    s = d.slide("Data analysis: what differs from the paper")
    rows = [("", "paper", "Binance"), ("spread of A_{C}, Q75/Q25", v["pvu.QratioAC.paper"], v["pvu.QratioAC"]),
            ("corr(A, |Ret|)", v["pvu.corrAabs.paper"], v["pvu.corrAabs"]),
            ("s.d. of monthly return", "—", v["d.retSd"] + "%")]
    tbl = s.shapes.add_table(len(rows), 3, Inches(0.47), Inches(1.75), Inches(5.2), Inches(1.8)).table
    for r in range(len(rows)):
        tbl.rows[r].height = Inches(0.45)
    tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(2.8), Inches(1.2), Inches(1.2)
    tpr = tbl._tbl.tblPr
    for attr in ["firstRow", "bandRow"]:
        tpr.set(attr, "0")
    style = tpr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.background()
            c.margin_left = c.margin_right = Inches(0.04)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            runs(p, val, 14, bold=(i == 0))
    for yy, hh in [(1.75, 0.02), (1.75 + 0.45, 0.01), (1.75 + 1.8, 0.02)]:
        bar(s, 0.47, yy, 5.2, hh, HSE)
    text(s, 6.2, 2.0, 3.4, 1.4, "Coins differ less in liquidity, move much more, and trade more when volatile.", 18,
         spacing=1.05)

    s = d.slide("One outlier month: March 2021")
    picture(s, base / "fig_hl_cumulative.png", y=1.4, h=3.3, center_x=5.0)
    d.sentence(s, "The altcoin season makes the volume spread. Without it: %s%% per month." % v["d.exMeanAC"], 4.85)

    s = d.slide("Our replication next to the paper")
    picture(s, base / "fig_quintiles.png", y=1.38, h=3.35, center_x=5.0)
    d.sentence(s, "Same signs, but noisy: H-L on A_{C} %s." % v["pvu.HLACrawCoef"], 4.85)

    s = d.slide("Upgrades")
    for x, val, lab, size in [(0.47, v["hf.corr.A"], "corr with 5-min λ\n(paper %s)" % v["hf.corr.A.paper"], 40),
                              (2.77, v["an.covneg"], "months with COV < 0\n(Amihud & Noh premise)", 40),
                              (5.07, v["se.ls"], "sentiment long-short", 22),
                              (7.37, v["mx.fm2"], "ln A_{C} on MOEX", 22)]:
        d.stat(s, x, 2.1, 2.2, val, lab, size)
    d.sentence(s, "Amihud tracks true price impact, but no upgrade settles the debate.", 4.25)

    s = d.slide("The paper versus us")
    for x, val, lab in [(0.6, "%s/%s" % (v["sc.samesign"], v["sc.ls_n"]), "signs agree"),
                        (3.6, "%s/%s" % (v["sc.rep196"], v["sc.paper_sig"]), "significant results repeat"),
                        (6.6, v["mde.hlAC.ratio"] + "×", "smallest premium we could\ndetect, vs the paper's")]:
        d.stat(s, x, 2.1, 2.9, val, lab, 48)
    d.sentence(s, "The direction replicates. With %s months instead of %s, significance does not."
               % (v["n.months"], v["d.paperMonths"]), 4.25)

    s = d.slide("Overfitting")
    for x, val, lab in [(0.6, "%s/%s" % (v["lg.base.196"], v["lg.base.n"]),
                         "tests pass 1.96\n(%s–%s expected by chance)" % (v["lg.base.exp"], v["w0.exp_base"])),
                        (3.6, v["lg.base.bonf"], "survive Bonferroni"),
                        (6.6, "%s/%s" % (v["rb.robust"], v["rb.n"]), "stable in the bootstrap")]:
        d.stat(s, x, 2.1, 2.9, val, lab, 48)
    d.sentence(s, "Every test logged: our evidence is weak, and we say so.", 4.25)

    s = d.slide("Takeaways")
    text(s, 0.47, 1.9, 9.0, 2.4, ["Amihud is mostly volume.", "The paper's signs replicate on crypto.",
                                  "Its significance needs more than %s months." % v["n.months"]], 28, spacing=1.5)

    s = d.slide(header=False)
    picture(s, HERE / "hse/ornament.png", h=5.625).left = int(Inches(10) - Inches(5.625) * 1920 / 1080)
    picture(s, HERE / "hse/logo_full.png", 0.47, 0.28, h=0.72)
    text(s, 1.4, 0.45, 7.5, 0.4, "Price Impact or Trading Volume (Lou, Shu, 2017)", 22, anchor=MSO_ANCHOR.MIDDLE)
    picture(s, HERE / "hse/meme.png", y=1.25, h=4.0, center_x=4.75)

    d.prs.save(out)
    return d.n


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--speech", default=None)
    ap.add_argument("--out", default=str(HERE / "paper15_team12_LouShu2017.pptx"))
    a = ap.parse_args()
    n = build(load_values(HERE / "gen/values.tex"), load_notes(a.speech), a.out)
    print(f"wrote {a.out} ({n} slides)")


if __name__ == "__main__":
    main()
